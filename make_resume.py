from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


OUT = "Mahrosh_Hamid_RBS_Resume.docx"


def set_cell_margins(cell, top=0, start=0, bottom=0, end=0):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tcMar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def hide_borders(table):
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        elem = OxmlElement(f"w:{edge}")
        elem.set(qn("w:val"), "nil")
        borders.append(elem)
    tblPr.append(borders)


def set_font(run, bold=None, size=None, italic=None):
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if size is not None:
        run.font.size = Pt(size)


def para(doc_or_cell, text="", alignment=WD_ALIGN_PARAGRAPH.LEFT, space_before=0, space_after=0, line_spacing=1.0):
    p = doc_or_cell.add_paragraph() if hasattr(doc_or_cell, "add_paragraph") else doc_or_cell.paragraphs[0]
    p.alignment = alignment
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = line_spacing
    if text:
        r = p.add_run(text)
        set_font(r, size=10)
    return p


def add_section(doc, title):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(8)
    pf.space_after = Pt(3)
    r = p.add_run(title.upper())
    set_font(r, bold=True, size=10.5)


def add_entry(doc, left, right, details=None):
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(4.95)
    table.columns[1].width = Inches(2.05)
    hide_borders(table)
    for cell in table.rows[0].cells:
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    l = table.cell(0, 0).paragraphs[0]
    l.paragraph_format.space_after = Pt(0)
    r = l.add_run(left)
    set_font(r, bold=True, size=10)
    rr = table.cell(0, 1).paragraphs[0]
    rr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    rr.paragraph_format.space_after = Pt(0)
    r = rr.add_run(right)
    set_font(r, size=10)
    if details:
        p = doc.add_paragraph(style=None)
        p.paragraph_format.left_indent = Inches(0.18)
        p.paragraph_format.first_line_indent = Inches(-0.18)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.0
        run = p.add_run("• " + details)
        set_font(run, size=9.5)


doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.42)
sec.bottom_margin = Inches(0.42)
sec.left_margin = Inches(0.55)
sec.right_margin = Inches(0.55)

styles = doc.styles
styles["Normal"].font.name = "Arial"
styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
styles["Normal"].font.size = Pt(10)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(1)
r = p.add_run("MAHROSH HAMID")
set_font(r, bold=True, size=16)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(5)
r = p.add_run("Robbinsville, NJ  |  609 977 3443  |  hamidmahrosh@gmail.com")
set_font(r, size=9.5)

add_section(doc, "Education")
add_entry(doc, "Rutgers University New Brunswick", "Current Student")
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(1)
r = p.add_run("Biomedical Engineering")
set_font(r, italic=True, size=9.5)
add_entry(doc, "Robbinsville High School", "September 2022 to June 2026")
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(1)
r = p.add_run("High School Diploma")
set_font(r, italic=True, size=9.5)

add_section(doc, "Experience")
add_entry(doc, "Muslim Center of Greater Princeton, Islamic Teacher", "September 2023 to Present", "Teach Islamic history and core principles to a class of 25 or more students ages seven and eight. Completed more than 400 hours of community service.")
add_entry(doc, "Biskidi Mendhi, Founder and Henna Artist", "2024 to Present", "Launched a henna business and provided custom designs at five events, earning more than $50 per event.")
add_entry(doc, "Muslim Student Association, Social Media and Outreach Coordinator", "September 2025 to June 2026", "Created promotional posters and social media posts for club events. Coordinated outreach with local Muslim Student Associations, sponsors, and caterers.")
add_entry(doc, "Luv Michael Nonprofit Company, Community Service Volunteer", "March 2025 to April 2025", "Supported an autism awareness campaign by contacting local representatives and researching resource centers. Raised more than $200 and completed 40 hours of community service.")
add_entry(doc, "Muslim Center of Greater Princeton Free Clinic, Refugee Tutor and Volunteer", "March 2025 to April 2025", "Collaborated with a clinical therapist and social work director to develop math and English lessons for Afghan refugee students from prekindergarten through eighth grade. Completed 40 hours of community service.")

add_section(doc, "Skills")
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(1)
r = p.add_run("Team Leadership  |  Strategic Thinking  |  Marketing  |  Team Engagement")
set_font(r, size=9.5)

add_section(doc, "Honors and Awards")
p = doc.add_paragraph()
p.paragraph_format.left_indent = Inches(0.18)
p.paragraph_format.first_line_indent = Inches(-0.18)
p.paragraph_format.space_after = Pt(1)
r = p.add_run("• Princeton University IgniteSTEM, second place winner, Urban Utopia Design Challenge, Fall 2024")
set_font(r, size=9.5)
p = doc.add_paragraph()
p.paragraph_format.left_indent = Inches(0.18)
p.paragraph_format.first_line_indent = Inches(-0.18)
p.paragraph_format.space_after = Pt(1)
r = p.add_run("• Member, World Language Honor Society and National Honor Society, Robbinsville High School")
set_font(r, size=9.5)

add_section(doc, "Languages")
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(0)
r = p.add_run("Urdu, native proficiency  |  English, native proficiency  |  Spanish")
set_font(r, size=9.5)

doc.save(OUT)
print(OUT)
